"""Work files (Source 1 format): validation on upload and the multi-episode download."""
import re
from copy import copy, deepcopy
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.drawing.image import Image

from app.services.rough_parser import _dur_to_sec, _s
from app.services.sheet_kind import (
    ROUGH_AS_WORK_FILE_MSG,
    SheetFormatError,
    has_rough_song_table,
    work_file_header_row,
)

# "ep-25", "Ep 25", "episode_25" — the episode marker used in work-file titles and filenames.
_EP_RE = re.compile(r"\bep(?:isode)?[\s._-]*(\d{1,4})", re.IGNORECASE)


def _ep_from_text(text: str) -> int | None:
    m = _EP_RE.search(text or "")
    return int(m.group(1)) if m else None


def _as_int(v) -> int | None:
    try:
        n = int(float(_s(v)))
    except ValueError:
        return None
    return n if n > 0 else None


def _norm(text: str) -> str:
    """Case, spacing and punctuation don't count: 'Lakshmi  niwas' == 'Lakshmi-Niwas'."""
    return re.sub(r"[^a-z0-9]", "", (text or "").lower())


def serial_from_title(title: str) -> str:
    """'Lakshmi niwas-ep-1-12 january 2026.mp4' → 'Lakshmi niwas' (whole title when there's no ep marker)."""
    m = _EP_RE.search(title or "")
    name = title[:m.start()] if m else re.sub(r"\.(mp4|mov|mkv|avi|mxf)$", "", title or "", flags=re.IGNORECASE)
    return name.strip(" -_.")


def _header_fields(rows: list) -> dict:
    """Label → value pairs from the block above the cue table (label cell, value in the next cell)."""
    wanted = {"title": "title", "episode number": "episode_number", "duration": "duration"}
    out: dict = {}
    for r in rows:
        for j, v in enumerate(r[:-1]):
            key = wanted.get(_s(v).lower())
            if key and key not in out and r[j + 1] is not None:
                out[key] = r[j + 1]
    return out


def _work_sheets(wb) -> tuple[list, bool]:
    """(work-file sheets as (ws, rows, header_idx), whether any rough-sheet tab is present)."""
    found, rough = [], False
    for ws in wb.worksheets:
        rows = [list(r) + [None] * max(0, 11 - len(r)) for r in ws.iter_rows(values_only=True)]
        hdr = work_file_header_row(rows)
        if hdr is not None:
            found.append((ws, rows, hdr))
        elif has_rough_song_table(rows):
            rough = True
    return found, rough


def inspect_work_file(data: bytes, filename: str, project_title: str) -> dict:
    """Validate one uploaded work file and extract what the upload preview shows.

    Raises SheetFormatError when the file is not a single-episode work file.
    """
    try:
        wb = load_workbook(BytesIO(data), data_only=True)
    except Exception:
        raise SheetFormatError("Couldn't read this file. Upload the work file as .xlsx.")

    found, rough = _work_sheets(wb)
    if not found:
        raise SheetFormatError(
            ROUGH_AS_WORK_FILE_MSG if rough else
            "Not a Work File: expected a 'Cue Sheet' with TITLE / CONTRIBUTORS / TIME IN / TIME OUT / DURATION columns."
        )
    if rough:
        raise SheetFormatError("This file also contains Rough Sheet tabs. Upload the work file on its own.")
    if len(found) > 1:
        raise SheetFormatError(f"This file has {len(found)} cue sheets. Upload one work file per episode.")

    _, rows, hdr = found[0]
    fields = _header_fields(rows[:hdr])
    title = _s(fields.get("title"))
    cue_count = sum(1 for r in rows[hdr + 1:] if _s(r[0]))
    if cue_count == 0:
        raise SheetFormatError("This work file has no cues under the TITLE / CONTRIBUTORS header.")

    warnings: list[str] = []

    # Episode number: explicit field → sheet title → filename; flag any disagreement.
    candidates = [
        ("Episode number field", _as_int(fields.get("episode_number"))),
        ("sheet title", _ep_from_text(title)),
        ("filename", _ep_from_text(filename)),
    ]
    detected = [(src, n) for src, n in candidates if n]
    episode_number = detected[0][1] if detected else None
    if len({n for _, n in detected}) > 1:
        warnings.append("Episode numbers disagree — " + ", ".join(f"{src}: Ep {n}" for src, n in detected))
    if episode_number is None:
        warnings.append("Episode number not found in the file — enter it before saving.")

    # Serial name check — a mismatch needs the uploader's confirmation before it can be saved.
    serial_in_file = serial_from_title(title)

    return {
        "episode_number": episode_number,
        "title": title,
        "cue_count": cue_count,
        "duration_sec": _dur_to_sec(fields.get("duration")) or None,
        "serial_in_file": serial_in_file,
        "serial_mismatch": _norm(serial_in_file) != _norm(project_title),
        "warnings": warnings,
    }


def _copy_sheet(src, dst) -> None:
    for row in src.iter_rows():
        for c in row:
            if isinstance(c, MergedCell):
                continue
            d = dst.cell(row=c.row, column=c.column, value=c.value)
            if c.has_style:
                d.font = copy(c.font)
                d.border = copy(c.border)
                d.fill = copy(c.fill)
                d.number_format = c.number_format
                d.protection = copy(c.protection)
                d.alignment = copy(c.alignment)
    for rng in src.merged_cells.ranges:
        dst.merge_cells(str(rng))
    for key, dim in src.column_dimensions.items():
        dst.column_dimensions[key].width = dim.width
        dst.column_dimensions[key].hidden = dim.hidden
    for idx, dim in src.row_dimensions.items():
        if dim.height:
            dst.row_dimensions[idx].height = dim.height
    dst.freeze_panes = src.freeze_panes
    dst.sheet_format = copy(src.sheet_format)
    dst.page_margins = copy(src.page_margins)
    dst.print_options = copy(src.print_options)
    for attr in ("orientation", "paperSize", "scale", "fitToWidth", "fitToHeight"):
        setattr(dst.page_setup, attr, getattr(src.page_setup, attr))
    # Embedded pictures (the cue-sheet logo). openpyxl only reads these when Pillow is installed.
    for img in src._images:
        pic = Image(BytesIO(img._data()))
        pic.width, pic.height = img.width, img.height
        pic.anchor = deepcopy(img.anchor)
        dst.add_image(pic)


def build_work_file_bulk(files: list) -> bytes:
    """One workbook, one sheet per episode (named by episode number, like the IPRS bulk export)."""
    out = Workbook()
    out.remove(out.active)
    for wf in sorted(files, key=lambda f: f.episode_number):
        wb = load_workbook(BytesIO(wf.content))
        found, _ = _work_sheets(wb)
        src = found[0][0] if found else wb.worksheets[0]
        _copy_sheet(src, out.create_sheet(title=str(wf.episode_number)[:31]))
    buf = BytesIO()
    out.save(buf)
    return buf.getvalue()
